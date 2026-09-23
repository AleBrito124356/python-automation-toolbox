#!/usr/bin/env python3
"""Convert between CSV and Excel (.xlsx) with delimiter and encoding sniffing.

Direction is inferred from the file extensions. CSV to xlsx sniffs the
delimiter (comma, semicolon, tab, pipe) and the encoding (UTF-8, falling
back to Windows-1252 for files exported by Excel in Spanish/European
locales), and converts cells that are plainly numbers to real numbers so
Excel treats them as such. Everything else stays text, exactly as written:

  - leading zeros stay text (ZIP codes, phone numbers, IDs: 07001, -007)
  - a leading + stays text (phone numbers like +50760000000)
  - integers longer than 15 digits stay text (Excel would round them)
  - words Python happens to parse as numbers stay text (Nan, Infinity, 1_000)
  - scientific notation (1e5, 12E3) stays text unless --allow-exponent
  - decimal commas (3,50 or 1.234,56) become numbers only with --decimal-comma

Xlsx to CSV writes UTF-8 with BOM so Excel reopens accented characters
correctly; --decimal-comma writes 3,5 and switches the delimiter to ';'.

Usage:
    python 08_csv_excel.py ventas.csv ventas.xlsx
    python 08_csv_excel.py ventas_es.csv ventas.xlsx --decimal-comma
    python 08_csv_excel.py report.xlsx report.csv --sheet "Q3 Data"
    python 08_csv_excel.py report.xlsx --list-sheets

Dependencies: openpyxl
"""

from __future__ import annotations

import argparse
import codecs
import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

INT_RE = re.compile(r"-?\d+")
FLOAT_RE = re.compile(r"-?(?:\d+\.\d+|\.\d+|\d+\.)")
EXP_RE = re.compile(r"-?(?:\d+\.?\d*|\.\d+)[eE][+-]?\d+")
# Decimal-comma locales: '.' groups thousands, ',' marks decimals.
THOUSANDS_INT_RE = re.compile(r"-?\d{1,3}(?:\.\d{3})+")
DECIMAL_COMMA_RE = re.compile(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d+")
MAX_EXACT_DIGITS = 15            # Excel keeps 15 significant digits
BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def coerce(value: str, *, decimal_comma: bool = False, allow_exponent: bool = False):
    """Return int/float for cells that are plainly numbers, else the original text."""
    s = value.strip()
    if s == "":
        return None
    unsigned = s[1:] if s.startswith("-") else s
    if len(unsigned) > 1 and unsigned[0] == "0" and unsigned[1] not in ".,":
        return value  # ZIP codes, phone numbers, IDs
    if INT_RE.fullmatch(s):
        if len(unsigned) > MAX_EXACT_DIGITS:
            return value  # card numbers, long IDs: Excel would round them
        return int(s)
    if decimal_comma:
        # 1.234 means one thousand two hundred thirty-four here, never 1.234.
        if THOUSANDS_INT_RE.fullmatch(s):
            return int(s.replace(".", ""))
        if DECIMAL_COMMA_RE.fullmatch(s):
            return float(s.replace(".", "").replace(",", "."))
        return value
    if FLOAT_RE.fullmatch(s):
        return float(s)
    if allow_exponent and EXP_RE.fullmatch(s):
        return float(s)
    return value


def detect_encoding(path: Path) -> str:
    """UTF-8 (with or without BOM) if the whole file decodes, else Windows-1252."""
    decoder = codecs.getincrementaldecoder("utf-8")()
    try:
        with path.open("rb") as fh:
            while chunk := fh.read(1024 * 1024):
                decoder.decode(chunk)
        decoder.decode(b"", final=True)
    except UnicodeDecodeError:
        return "cp1252"
    return "utf-8-sig"


def sniff_delimiter(path: Path, encoding: str) -> str:
    with path.open("r", encoding=encoding, newline="") as fh:
        sample = fh.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def sheet_title(stem: str) -> str:
    title = BAD_SHEET_CHARS.sub("_", stem).strip("'")[:31]  # Excel's sheet-name rules
    return title or "Sheet1"


def csv_to_xlsx(src: Path, dest: Path, *, encoding: str = "auto",
                decimal_comma: bool = False, allow_exponent: bool = False) -> int:
    if encoding == "auto":
        encoding = detect_encoding(src)
        if encoding != "utf-8-sig":
            print(f"Not valid UTF-8; reading as {encoding} (override with --encoding).")
    delimiter = sniff_delimiter(src, encoding)
    shown = {"\t": "TAB"}.get(delimiter, delimiter)
    print(f"Detected delimiter: '{shown}'")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title(src.stem)
    rows = 0
    with src.open("r", encoding=encoding, newline="") as fh:
        for row in csv.reader(fh, delimiter=delimiter):
            ws.append([coerce(cell, decimal_comma=decimal_comma, allow_exponent=allow_exponent)
                       for cell in row])
            rows += 1
    ws.freeze_panes = "A2"
    wb.save(dest)
    print(f"Wrote {rows} rows -> {dest}")
    return rows


def format_cell(cell, decimal_comma: bool) -> object:
    if cell is None:
        return ""
    if decimal_comma and isinstance(cell, float):
        return repr(cell).replace(".", ",")
    return cell


def xlsx_to_csv(src: Path, dest: Path, sheet: str | None, delimiter: str | None,
                decimal_comma: bool = False) -> int:
    if delimiter is None:
        delimiter = ";" if decimal_comma else ","
    if decimal_comma and delimiter == ",":
        sys.exit("--decimal-comma needs a delimiter other than ',' (default is ';').")
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if sheet is None:
        ws = wb.active
    elif sheet in wb.sheetnames:
        ws = wb[sheet]
    elif sheet.isdigit() and int(sheet) < len(wb.sheetnames):
        ws = wb[wb.sheetnames[int(sheet)]]
    else:
        sys.exit(f"Sheet '{sheet}' not found. Available: {', '.join(wb.sheetnames)}")

    rows = 0
    with dest.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([format_cell(cell, decimal_comma) for cell in row])
            rows += 1
    wb.close()
    print(f"Wrote {rows} rows from sheet '{ws.title}' -> {dest}")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="CSV <-> xlsx converter with sniffing.")
    parser.add_argument("input", help="source file (.csv or .xlsx)")
    parser.add_argument("output", nargs="?", help="destination file (.xlsx or .csv)")
    parser.add_argument("--sheet", help="xlsx->csv: sheet name or 0-based index (default: active)")
    parser.add_argument("--delimiter",
                        help="xlsx->csv: output delimiter (default: ',' or ';' with --decimal-comma)")
    parser.add_argument("--decimal-comma", action="store_true",
                        help="numbers use a decimal comma (3,50 / 1.234,56), as in es/de/fr locales")
    parser.add_argument("--allow-exponent", action="store_true",
                        help="csv->xlsx: treat 1e5 / 12E3 as numbers (default: keep as text)")
    parser.add_argument("--encoding", default="auto",
                        help="csv->xlsx: input encoding (default: auto = UTF-8, else cp1252)")
    parser.add_argument("--list-sheets", action="store_true", help="print sheet names and exit")
    args = parser.parse_args()

    src = Path(args.input).expanduser()
    if not src.is_file():
        sys.exit(f"Not found: {src}")

    if args.list_sheets:
        if src.suffix.lower() != ".xlsx":
            sys.exit("--list-sheets only applies to .xlsx files")
        wb = openpyxl.load_workbook(src, read_only=True)
        for i, name in enumerate(wb.sheetnames):
            print(f"  [{i}] {name}")
        return

    if not args.output:
        sys.exit("Provide an output file (or use --list-sheets).")
    dest = Path(args.output).expanduser()

    pair = (src.suffix.lower(), dest.suffix.lower())
    if pair == (".csv", ".xlsx"):
        try:
            csv_to_xlsx(src, dest, encoding=args.encoding, decimal_comma=args.decimal_comma,
                        allow_exponent=args.allow_exponent)
        except (LookupError, UnicodeDecodeError) as exc:
            sys.exit(f"Could not read {src.name} as {args.encoding}: {exc}")
    elif pair == (".xlsx", ".csv"):
        xlsx_to_csv(src, dest, args.sheet, args.delimiter, args.decimal_comma)
    else:
        sys.exit("Expected csv->xlsx or xlsx->csv (check the file extensions).")


if __name__ == "__main__":
    main()
