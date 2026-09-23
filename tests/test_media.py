"""Media and document tools: 05 images, 06 PDFs, 07 video->GIF, 08 CSV/xlsx."""

from __future__ import annotations

import shutil
import subprocess

import openpyxl
import pytest
from PIL import Image
from pypdf import PdfReader, PdfWriter

from conftest import load_script

ORIENTATION = 0x0112


# --------------------------------------------------------------------------- 05 images

@pytest.fixture
def photos(tmp_path):
    src = tmp_path / "in"
    src.mkdir()
    Image.new("RGB", (1600, 800), (200, 30, 30)).save(src / "photo.jpg", quality=95)
    exif = Image.Exif()
    exif[ORIENTATION] = 6          # "rotate 90 CW to display"
    exif[0x010F] = "TestCam"       # Make
    Image.new("RGB", (2000, 1000), (30, 200, 30)).save(src / "rotated.jpg", exif=exif.tobytes())
    Image.new("RGBA", (300, 300), (0, 0, 255, 0)).save(src / "logo.png")
    Image.new("RGB", (300, 300), (0, 0, 255)).save(src / "logo.jpg")
    return src


def test_images_keep_real_format_and_quality(photos, tmp_path, run):
    """Audit repro: every JPEG came out as PNG data named .jpg, ignoring --quality."""
    out = tmp_path / "out"
    result = run("05", photos, "--out", out, "--resize", "800", "--quality", "60")
    assert result.returncode == 0, result.stderr
    for name, size in {"photo.jpg": (800, 400), "logo.jpg": (300, 300)}.items():
        with Image.open(out / name) as img:
            assert img.format == "JPEG"
            assert img.size == size
    with Image.open(out / "logo.png") as img:
        assert img.format == "PNG" and img.mode == "RGBA"
    low = (out / "photo.jpg").stat().st_size
    run("05", photos, "--out", tmp_path / "hq", "--resize", "800", "--quality", "98")
    assert (tmp_path / "hq" / "photo.jpg").stat().st_size > low, "--quality must matter"


def test_images_bake_orientation_and_reset_tag(photos, tmp_path, run):
    """Audit repro: pixels were rotated but Orientation=6 was kept -> rotated twice."""
    out = tmp_path / "out"
    run("05", photos, "--out", out, "--format", "jpeg")
    with Image.open(out / "rotated.jpg") as img:
        assert img.size == (1000, 2000)
        exif = img.getexif()
        assert exif.get(ORIENTATION) == 1
        assert exif.get(0x010F) == "TestCam", "other EXIF survives when not stripped"


def test_images_strip_exif_still_upright(photos, tmp_path, run):
    out = tmp_path / "out"
    run("05", photos, "--out", out, "--strip-exif")
    with Image.open(out / "rotated.jpg") as img:
        assert img.size == (1000, 2000)
        assert len(img.getexif()) == 0


def test_images_name_collisions_are_disambiguated(photos, tmp_path, run):
    """Audit repro: logo.png and logo.jpg both became logo.webp (one silently lost)."""
    out = tmp_path / "out"
    result = run("05", photos, "--out", out, "--format", "webp")
    assert result.returncode == 0, result.stderr
    assert sorted(p.name for p in out.iterdir()) == [
        "logo.jpg.webp", "logo.png.webp", "photo.webp", "rotated.webp"]
    assert "Processed 4/4" in result.stdout


def test_images_transparent_to_jpeg_is_white_not_black(photos, tmp_path, run):
    out = tmp_path / "out"
    run("05", photos, "--out", out, "--format", "jpeg")
    with Image.open(out / "logo.png.jpg") as img:
        assert img.getpixel((10, 10)) == (255, 255, 255)


def test_images_16bit_png_resizes_and_converts(tmp_path, run):
    src = tmp_path / "scans"
    src.mkdir()
    Image.new("I;16", (400, 200), 51200).save(src / "depth.png")
    result = run("05", src, "--out", tmp_path / "o", "--resize", "100")
    assert result.returncode == 0, result.stderr
    with Image.open(tmp_path / "o" / "depth.png") as img:
        assert (img.mode, img.size, img.getpixel((3, 3))) == ("I;16", (100, 50), 51200)
    run("05", src, "--out", tmp_path / "j", "--format", "jpeg")
    with Image.open(tmp_path / "j" / "depth.jpg") as img:
        assert img.mode == "L" and abs(img.getpixel((3, 3)) - 200) <= 1   # scaled, not clipped


def test_images_bad_resize_is_an_argparse_error(photos, tmp_path, run):
    result = run("05", photos, "--out", tmp_path / "o", "--resize", "big")
    assert result.returncode == 2
    assert "not a size" in result.stderr


# --------------------------------------------------------------------------- 06 PDFs

@pytest.fixture
def pdfs(tmp_path):
    def make(name: str, pages: int, width: int = 200):
        writer = PdfWriter()
        for i in range(pages):
            writer.add_blank_page(width=width + i, height=300)
        path = tmp_path / name
        with path.open("wb") as fh:
            writer.write(fh)
        return path
    return make


def page_widths(path) -> list[int]:
    return [int(p.mediabox.width) for p in PdfReader(path).pages]


def test_pdf_merge_split_rotate_extract(pdfs, tmp_path, run):
    a, b = pdfs("a.pdf", 2, 100), pdfs("b.pdf", 3, 200)
    assert run("06", "merge", a, b, "-o", tmp_path / "both.pdf").returncode == 0
    assert page_widths(tmp_path / "both.pdf") == [100, 101, 200, 201, 202]

    assert run("06", "split", tmp_path / "both.pdf", "-o", tmp_path / "pages").returncode == 0
    assert len(list((tmp_path / "pages").glob("both_p*.pdf"))) == 5

    assert run("06", "rotate", tmp_path / "both.pdf", "--angle", "90", "--pages", "2-3",
               "-o", tmp_path / "rot.pdf").returncode == 0
    assert [p.rotation for p in PdfReader(tmp_path / "rot.pdf").pages] == [0, 90, 90, 0, 0]

    assert run("06", "extract", tmp_path / "both.pdf", "--pages", "4-",
               "-o", tmp_path / "tail.pdf").returncode == 0
    assert page_widths(tmp_path / "tail.pdf") == [201, 202]
    assert run("06", "extract", tmp_path / "both.pdf", "--pages", "-2,5",
               "-o", tmp_path / "mix.pdf").returncode == 0
    assert page_widths(tmp_path / "mix.pdf") == [100, 101, 202]


@pytest.mark.parametrize("spec, message", [
    ("3-", "out of bounds"),
    ("abc", "is not a page or range"),
    ("1,,2", "empty item"),
    ("2-1", "out of bounds"),
])
def test_pdf_bad_page_specs_are_one_line_errors(pdfs, tmp_path, run, spec, message):
    """Audit repro: '--pages 3-' and '--pages abc' raised raw ValueError tracebacks."""
    src = pdfs("two.pdf", 2)
    result = run("06", "extract", src, "--pages", spec, "-o", tmp_path / "x.pdf")
    assert result.returncode == 1
    assert message in result.stderr
    assert "Traceback" not in result.stderr


def test_pdf_refuses_to_overwrite_its_input(pdfs, run):
    src = pdfs("scan.pdf", 2)
    before = src.read_bytes()
    result = run("06", "rotate", src, "--angle", "90", "-o", src)
    assert result.returncode == 1 and "also an input" in result.stderr
    assert src.read_bytes() == before


def test_pdf_password_protected_input(tmp_path, run):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt(user_password="s3cret", algorithm="RC4-128")
    locked = tmp_path / "locked.pdf"
    with locked.open("wb") as fh:
        writer.write(fh)
    denied = run("06", "split", locked, "-o", tmp_path / "p")
    assert denied.returncode == 1 and "--password" in denied.stderr
    allowed = run("06", "split", locked, "-o", tmp_path / "p", "--password", "s3cret")
    assert allowed.returncode == 0, allowed.stderr


# --------------------------------------------------------------------------- 07 video -> GIF

@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="ffmpeg not installed")
def test_video_to_gif_two_pass(tmp_path, run):
    clip = tmp_path / "clip.mp4"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi",
                    "-i", "testsrc=duration=1:size=320x240:rate=10", "-pix_fmt", "yuv420p",
                    str(clip)], check=True, timeout=60)
    result = run("07", clip, "--fps", "5", "--width", "160")
    assert result.returncode == 0, result.stderr
    with Image.open(tmp_path / "clip.gif") as gif:
        assert gif.format == "GIF"
        assert gif.size[0] == 160
        assert gif.n_frames >= 4


def test_video_to_gif_missing_input(tmp_path, run):
    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg not installed")
    result = run("07", tmp_path / "nope.mp4")
    assert result.returncode == 1 and "Not found" in result.stderr


# --------------------------------------------------------------------------- 08 CSV <-> xlsx

def test_csv_roundtrip_keeps_words_and_codes_intact(tmp_path, run):
    """Audit repro: 'Nan'/'Infinity' vanished, '1e5'/'1_000' became numbers."""
    src = tmp_path / "people.csv"
    src.write_text("name;zip;code;qty;price\n"
                   "Nan;07001;1e5;1_000;3,50\n"
                   "Infinity;90210;12E3;7;2.5\n"
                   "-007;+50760000000;1234567890123456789;-4;0.25\n", encoding="utf-8")
    assert run("08", src, tmp_path / "people.xlsx").returncode == 0
    rows = list(openpyxl.load_workbook(tmp_path / "people.xlsx").active.iter_rows(values_only=True))
    assert rows[1] == ("Nan", "07001", "1e5", "1_000", "3,50")
    assert rows[2] == ("Infinity", 90210, "12E3", 7, 2.5)
    assert rows[3] == ("-007", "+50760000000", "1234567890123456789", -4, 0.25)

    assert run("08", tmp_path / "people.xlsx", tmp_path / "back.csv").returncode == 0
    back = (tmp_path / "back.csv").read_text(encoding="utf-8-sig").splitlines()
    assert back[1] == 'Nan,07001,1e5,1_000,"3,50"'
    assert back[2] == "Infinity,90210,12E3,7,2.5"


def test_csv_opt_in_number_formats():
    mod = load_script("08")
    assert mod.coerce("1e5") == "1e5"
    assert mod.coerce("1e5", allow_exponent=True) == 100000.0
    assert mod.coerce("3,50", decimal_comma=True) == 3.5
    assert mod.coerce("1.234,56", decimal_comma=True) == 1234.56
    assert mod.coerce("1.234", decimal_comma=True) == 1234
    assert mod.coerce("0,5", decimal_comma=True) == 0.5
    assert mod.coerce("3,50") == "3,50"
    assert mod.coerce("0") == 0 and mod.coerce("0.5") == 0.5
    assert mod.coerce("") is None


def test_csv_decimal_comma_roundtrip_and_cp1252(tmp_path, run):
    src = tmp_path / "ventas.csv"
    src.write_bytes("producto;precio;unidades\nCafé;3,50;1.200\n".encode("cp1252"))
    result = run("08", src, tmp_path / "v.xlsx", "--decimal-comma")
    assert result.returncode == 0, result.stderr
    assert "cp1252" in result.stdout
    rows = list(openpyxl.load_workbook(tmp_path / "v.xlsx").active.iter_rows(values_only=True))
    assert rows[1] == ("Café", 3.5, 1200)
    assert run("08", tmp_path / "v.xlsx", tmp_path / "v.csv", "--decimal-comma").returncode == 0
    assert (tmp_path / "v.csv").read_text(encoding="utf-8-sig").splitlines()[1] == "Café;3,5;1200"


def test_xlsx_sheet_selection_and_listing(tmp_path, run):
    wb = openpyxl.Workbook()
    wb.active.title = "Resumen"
    wb.active.append(["a"])
    second = wb.create_sheet("Q3 Data")
    second.append(["x", 1])
    wb.save(tmp_path / "r.xlsx")
    listing = run("08", tmp_path / "r.xlsx", "--list-sheets")
    assert "[1] Q3 Data" in listing.stdout
    assert run("08", tmp_path / "r.xlsx", tmp_path / "q3.csv", "--sheet", "Q3 Data").returncode == 0
    assert (tmp_path / "q3.csv").read_text(encoding="utf-8-sig").strip() == "x,1"
    missing = run("08", tmp_path / "r.xlsx", tmp_path / "n.csv", "--sheet", "Nope")
    assert missing.returncode == 1 and "Available: Resumen, Q3 Data" in missing.stderr
